"""Pruebas del análisis de eficiencia en tokens."""

import unittest
from tempfile import TemporaryDirectory
from unittest.mock import Mock
from pathlib import Path

from openpyxl import Workbook, load_workbook

from metricas.tokens import (
    analizar_tokens,
    clasificar_eficiencia,
    crear_metrica_tokens,
)


class TokensTest(unittest.TestCase):
    """Verifica cálculos, límites y validaciones de tokens."""

    def test_calcula_eficiencias_y_penaliza_refinamientos(self):
        """Calcula NLOC por mil tokens y aplica el coeficiente."""
        valores = {"Método utilizado": "Chat", "Refinamientos": 2,
                   "Tokens registrados": 5000}
        metrica = crear_metrica_tokens("M1", valores, 100, 0.1)
        self.assertEqual(metrica.eficiencia_generacion, 20)
        self.assertEqual(metrica.eficiencia_ponderada, 16.67)
        self.assertEqual(metrica.nivel_eficiencia, "Medio")

    def test_clasificacion_respeta_limites(self):
        """Clasifica correctamente los valores 10 y 30."""
        self.assertEqual(clasificar_eficiencia(10)[0], "Bajo")
        self.assertEqual(clasificar_eficiencia(10.0001)[0], "Medio")
        self.assertEqual(clasificar_eficiencia(29.9999)[0], "Medio")
        self.assertEqual(clasificar_eficiencia(30)[0], "Alto")

    def test_clasifica_antes_de_redondear(self):
        """Usa el valor original aunque el almacenado se redondee a 30."""
        valores = {"Método utilizado": "IDE", "Refinamientos": 0,
                   "Tokens registrados": 1000}
        metrica = crear_metrica_tokens("M1", valores, 29.999, 0.1)
        self.assertEqual(metrica.eficiencia_ponderada, 30)
        self.assertEqual(metrica.nivel_eficiencia, "Medio")

    def test_rechaza_datos_invalidos(self):
        """Rechaza método vacío, tokens no positivos y refinamientos decimales."""
        casos = [
            {"Método utilizado": "", "Refinamientos": 0, "Tokens registrados": 1},
            {"Método utilizado": "Chat", "Refinamientos": 0, "Tokens registrados": 0},
            {"Método utilizado": "Chat", "Refinamientos": 1.5, "Tokens registrados": 1},
        ]
        for valores in casos:
            with self.assertRaises(ValueError):
                crear_metrica_tokens("M1", valores, 10, 0.1)

    def test_lee_columnas_por_encabezado(self):
        """Lee correctamente aunque Código no sea la primera columna."""
        libro = Workbook()
        hoja = libro.active
        hoja.title = "Datos Tokens"
        hoja.append(["Tokens registrados", "Código", "Refinamientos",
                     "Método utilizado"])
        hoja.append([1000, "M1", 1, "Chat"])
        metrica = analizar_tokens(Mock(codigo="M1"), libro, 20, 0.1, Mock())
        self.assertEqual(metrica.metodo_utilizado, "Chat")

    def test_rechaza_codigo_duplicado(self):
        """No elige arbitrariamente entre dos filas del mismo proyecto."""
        libro = Workbook()
        hoja = libro.active
        hoja.title = "Datos Tokens"
        hoja.append(["Código", "Método utilizado", "Refinamientos",
                     "Tokens registrados"])
        hoja.append(["M1", "Chat", 0, 1000])
        hoja.append(["M1", "IDE", 1, 2000])
        with self.assertRaisesRegex(ValueError, "duplicado"):
            analizar_tokens(Mock(codigo="M1"), libro, 20, 0.1, Mock())

    def test_lee_libro_en_modo_solo_lectura(self):
        """No depende de max_row al buscar códigos en modo de solo lectura."""
        with TemporaryDirectory() as directorio:
            ruta = Path(directorio) / "entrada.xlsx"
            libro = Workbook()
            hoja = libro.active
            hoja.title = "Datos Tokens"
            hoja.append(["Código", "Método utilizado", "Refinamientos",
                         "Tokens registrados"])
            hoja.append(["M1", "Chat", 1, 1000])
            libro.save(ruta)
            libro.close()
            lectura = load_workbook(ruta, read_only=True, data_only=True)
            metrica = analizar_tokens(Mock(codigo="M1"), lectura, 20, 0.1, Mock())
            self.assertEqual(metrica.codigo, "M1")
            lectura.close()


if __name__ == "__main__":
    unittest.main()
